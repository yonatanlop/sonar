from __future__ import annotations

from io import BytesIO
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.shared.deps import get_db, get_current_user, require_analyst
from app.shared.exceptions import NotFoundError
from app.modules.actor_map.domain.actor import Actor, ActorRelation, ActorGroup
from app.modules.actor_map.infrastructure.actor_repository import (
    SqlActorRepository, SqlActorRelationRepository, SqlActorGroupRepository,
)

router = APIRouter(prefix="/actor-map", tags=["Actor Map"])


# ── Schemas ───────────────────────────────────────────────────
class ActorCreate(BaseModel):
    name: str
    actor_type: str
    party: str | None = None
    position: str | None = None
    country_code: str | None = None
    entity_id: UUID | None = None

class ActorOut(BaseModel):
    id: UUID
    name: str
    actor_type: str
    party: str | None
    position: str | None
    country_code: str | None
    entity_id: UUID | None
    active: bool

class ActorPatch(BaseModel):
    name: str | None = None
    actor_type: str | None = None
    party: str | None = None
    position: str | None = None
    active: bool | None = None

class RelationCreate(BaseModel):
    source_id: UUID
    target_id: UUID
    relation_type: str
    strength: int = 1

class RelationOut(BaseModel):
    id: UUID
    source_id: UUID
    target_id: UUID
    relation_type: str
    strength: int

class GroupCreate(BaseModel):
    name: str
    group_type: str
    color: str = "#6366f1"

class GroupOut(BaseModel):
    id: UUID
    name: str
    group_type: str
    color: str


# ── Actors ────────────────────────────────────────────────────
@router.get("/actors", response_model=list[ActorOut])
def list_actors(
    active_only: bool = True,
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    return [_actor_out(a) for a in SqlActorRepository(db).list(active_only=active_only)]


@router.post("/actors", response_model=ActorOut, status_code=201)
def create_actor(
    data: ActorCreate,
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    actor = Actor(
        id=uuid4(), name=data.name, actor_type=data.actor_type,
        party=data.party, position=data.position,
        country_code=data.country_code, entity_id=data.entity_id, active=True,
    )
    return _actor_out(SqlActorRepository(db).create(actor))


@router.patch("/actors/{actor_id}", response_model=ActorOut)
def update_actor(
    actor_id: UUID,
    data: ActorPatch,
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    repo = SqlActorRepository(db)
    actor = repo.get_by_id(actor_id)
    if not actor:
        raise NotFoundError("Actor no encontrado")
    if data.name is not None:
        actor.name = data.name
    if data.actor_type is not None:
        actor.actor_type = data.actor_type
    if data.party is not None:
        actor.party = data.party
    if data.position is not None:
        actor.position = data.position
    if data.active is not None:
        actor.active = data.active
    return _actor_out(repo.update(actor))


# ── Graph ─────────────────────────────────────────────────────
@router.get("/graph")
def get_graph(db: Session = Depends(get_db), _=Depends(get_current_user)):
    actors = SqlActorRepository(db).list(active_only=True)
    relations = SqlActorRelationRepository(db).list_all()
    nodes = [
        {
            "id": str(a.id),
            "data": {
                "label": a.name,
                "actor_type": a.actor_type,
                "party": a.party,
                "position": a.position,
            },
            "type": "actorNode",
        }
        for a in actors
    ]
    edges = [
        {
            "id": str(r.id),
            "source": str(r.source_id),
            "target": str(r.target_id),
            "label": r.relation_type,
            "data": {"strength": r.strength, "relation_type": r.relation_type},
        }
        for r in relations
    ]
    return {"nodes": nodes, "edges": edges}


# ── Relations ─────────────────────────────────────────────────
@router.post("/relations", response_model=RelationOut, status_code=201)
def create_relation(
    data: RelationCreate,
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    relation = ActorRelation(
        id=uuid4(), source_id=data.source_id, target_id=data.target_id,
        relation_type=data.relation_type, strength=data.strength,
    )
    saved = SqlActorRelationRepository(db).create(relation)
    return RelationOut(
        id=saved.id, source_id=saved.source_id, target_id=saved.target_id,
        relation_type=saved.relation_type, strength=saved.strength,
    )


@router.delete("/relations/{relation_id}", status_code=204)
def delete_relation(
    relation_id: UUID,
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    SqlActorRelationRepository(db).delete(relation_id)


# ── Groups ────────────────────────────────────────────────────
@router.get("/groups", response_model=list[GroupOut])
def list_groups(db: Session = Depends(get_db), _=Depends(get_current_user)):
    return [
        GroupOut(id=g.id, name=g.name, group_type=g.group_type, color=g.color)
        for g in SqlActorGroupRepository(db).list()
    ]


@router.post("/groups", response_model=GroupOut, status_code=201)
def create_group(
    data: GroupCreate,
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    group = ActorGroup(id=uuid4(), name=data.name, group_type=data.group_type, color=data.color)
    saved = SqlActorGroupRepository(db).create(group)
    return GroupOut(id=saved.id, name=saved.name, group_type=saved.group_type, color=saved.color)


# ── Excel Import ──────────────────────────────────────────────
@router.post("/import-excel", status_code=201)
async def import_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado")

    content = await file.read()
    wb = openpyxl.load_workbook(BytesIO(content))
    actor_repo = SqlActorRepository(db)
    relation_repo = SqlActorRelationRepository(db)

    created_actors = 0
    created_relations = 0

    if "Actores" in wb.sheetnames:
        ws = wb["Actores"]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            data = dict(zip(headers, row))
            actor = Actor(
                id=uuid4(),
                name=str(data.get("name", data.get("nombre", ""))),
                actor_type=str(data.get("actor_type", data.get("tipo", "politician"))),
                party=data.get("party", data.get("partido")),
                position=data.get("position", data.get("cargo")),
                country_code=data.get("country_code", data.get("pais")),
                entity_id=None,
                active=True,
            )
            actor_repo.create(actor)
            created_actors += 1

    if "Relaciones" in wb.sheetnames:
        ws = wb["Relaciones"]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        actors_by_name = {a.name: a for a in actor_repo.list(active_only=False)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            data = dict(zip(headers, row))
            src_name = str(data.get("source", data.get("origen", "")))
            tgt_name = str(data.get("target", data.get("destino", "")))
            src = actors_by_name.get(src_name)
            tgt = actors_by_name.get(tgt_name)
            if not src or not tgt:
                continue
            relation = ActorRelation(
                id=uuid4(), source_id=src.id, target_id=tgt.id,
                relation_type=str(data.get("relation_type", data.get("tipo_relacion", "ally"))),
                strength=int(data.get("strength", data.get("fuerza", 1))),
            )
            try:
                relation_repo.create(relation)
                created_relations += 1
            except Exception:
                db.rollback()

    return {"created_actors": created_actors, "created_relations": created_relations}


# ── Helper ────────────────────────────────────────────────────
def _actor_out(a: Actor) -> ActorOut:
    return ActorOut(
        id=a.id, name=a.name, actor_type=a.actor_type, party=a.party,
        position=a.position, country_code=a.country_code,
        entity_id=a.entity_id, active=a.active,
    )
